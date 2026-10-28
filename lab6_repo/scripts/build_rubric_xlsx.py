"""docs/cmap_rubric.xlsx: рубрика Новака, оцінки двох оцінювачів (формули), узгодженість, зауваження, precision/recall."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
ev = json.loads((ROOT / "docs/evidence/map_eval.json").read_text(encoding="utf-8"))
tr = [json.loads(l) for l in (ROOT / "models/triples_llm.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
rub = [l for l in (ROOT / "llm_log/prompts/step5_rubric_text.md").read_text(encoding="utf-8").splitlines() if l.startswith("| D")]
HEAD = PatternFill("solid", fgColor="D9E2F3"); BOLD = Font(bold=True); WRAP = Alignment(wrap_text=True, vertical="top")
wb = Workbook()


def sheet(name, rows, widths):
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.fill, c.font = HEAD, BOLD
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = WRAP
    return ws


wb.remove(wb.active)
# 1. Рубрика
rows = [["Критерій", "0", "1", "2", "3", "4"]]
for l in rub:
    rows.append([x.strip() for x in l.strip("|").split("|")])
sheet("Рубрика", rows, [40, 12, 22, 22, 22, 26])
# 2. Оцінки
MAPS = {"expert": "Власна (експертна v1)", "llm": "LLM-карта", "peer": "Карта «колеги» (LLM-замінник)"}
D = ["D1", "D2", "D3", "D4", "D5"]
rows = [["Карта", "Критерій", "Оцінювач 1", "Оцінювач 2", "Людина-оцінювач (заповнити)", "|різниця| (1–2)", "Середнє (1–2)"]]
r1, r2 = ev["scores"]["R1"], ev["scores"]["R2"]
for m, label in MAPS.items():
    for d in D:
        n = len(rows) + 1
        rows.append([label, d, r1[m][d], r2[m][d], None, f"=ABS(C{n}-D{n})", f"=AVERAGE(C{n}:D{n})"])
ws = sheet("Оцінки", rows, [32, 10, 14, 14, 26, 16, 16])
n = len(rows) + 2
ws.cell(n, 1, "Сума оцінювач 1 / 2").font = BOLD
for k, m in enumerate(MAPS):
    ws.cell(n + k, 1, f"Сума — {MAPS[m]}")
    ws.cell(n + k, 3, f"=SUM(C{2+5*k}:C{6+5*k})"); ws.cell(n + k, 4, f"=SUM(D{2+5*k}:D{6+5*k})")
# 3. Узгодженість
a = ev["agreement"]
last = 1 + 15
rows = [["Показник", "Формула в Excel", "Значення (Python)", "Інтерпретація"],
        ["Середня |різниця|", f"=AVERAGE(Оцінки!F2:F{last})", round(a["mean_abs_diff"], 3), "У середньому оцінювачі розходяться на 0,4 бала зі шкали 0–4: розбіжності малі."],
        ["Точний збіг", f'=COUNTIF(Оцінки!F2:F{last},0)/15', round(a["exact_agreement"], 3), "Дві третини оцінок ідентичні."],
        ["Збіг в межах 1 бала", f'=COUNTIF(Оцінки!F2:F{last},"<=1")/15', round(a["within_1"], 3), "Лише одна пара з 15 розійшлась більше ніж на 1 бал."],
        ["Кореляція Пірсона", f"=CORREL(Оцінки!C2:C{last},Оцінки!D2:D{last})", round(a["pearson"], 3), "Сильна позитивна кореляція (0,6–0,8): оцінювачі однаково впорядковують критерії."],
        ["Кореляція Спірмена", "—", round(a["spearman"], 3), "Узгоджена за рангами."],
        ["Квадратичний κ Коена", "—", round(a["weighted_kappa_quadratic"], 3), "«Суттєва» узгодженість (0,61–0,80) за шкалою Лендіса–Коха."],
        ["Розбіжність по критеріях", "—", "D2 = 1,0; D3 = 1,0; інші 0,0", "Найбільше сперечаються про глибину ієрархії та якість крос-зв'язків — критерії залежать від тлумачення «рівня» й «нетривіальності»."],
        ["Ранги карт", "—", "R1: peer > expert > llm; R2: expert > peer > llm", "Обидва оцінювачі ставлять LLM-карту останньою (немає крос-зв'язків і прикладів)."],
        ["Обмеження", "—", "—", "Обидва оцінювачі — незалежні запуски однієї моделі без доступу до відповідей одне одного; це не люди. Колонку «Людина-оцінювач» слід заповнити реальним рецензентом."]]
sheet("Узгодженість", rows, [28, 34, 26, 80])
# 4. Зауваження
rows = [["Карта", "№", "Зауваження (оцінювач)", "Враховано у v2"]]
FIX = {"expert": ["V2-01, V2-02", "V2-03", "V2-04, V2-05, V2-06"], "llm": ["V2-07", "V2-03, V2-06", "V2-08"], "peer": ["V2-08", "V2-06", "V2-07"]}
for m in MAPS:
    rem = ev["scores"]["R1" if m != "expert" else "R2"][m]["remarks"]
    for i, t in enumerate(rem[:3]):
        rows.append([MAPS[m], i + 1, t, FIX[m][i]])
sheet("Зауваження", rows, [28, 5, 110, 18])
# 5. Precision/recall
ex = {tuple(x) for x in ev["matches"]}
rows = [["ID", "Суб'єкт", "Відношення", "Об'єкт", "Є в експертній карті (strict)", "Текстово підтверджено (цитата + ручний перегляд)", "Помилки перевірки"]]
bad = {x[0] if isinstance(x, list) else x for x in ev["llm_text_supported_precision"]["auto_errors"] + ev["llm_text_supported_precision"]["manual_flags"]}
for t in tr:
    k = (t["s"].lower(), t["p"], t["o"].lower())
    rows.append([t["id"], t["s"], t["p"], t["o"], 1 if k in ex else 0, 0 if t["id"] in bad or t["checks"] else 1, ", ".join(str(c) for c in t.get("checks", []))])
ws = sheet("Precision-Recall", rows, [8, 32, 18, 32, 18, 24, 30])
nn = len(rows)
g = ev["llm_vs_expert_strict"]["gold"]
for i, (lab, f) in enumerate([("TP", f"=SUM(E2:E{nn})"), ("Передбачено (LLM)", f"=COUNTA(A2:A{nn})"), ("Еталон (експерт v1)", g), ("Precision", f"=B{nn+2}/B{nn+3}"), ("Recall", f"=B{nn+2}/B{nn+4}"), ("F1", f"=2*B{nn+5}*B{nn+6}/(B{nn+5}+B{nn+6})"), ("Text-supported precision", f"=SUM(F2:F{nn})/COUNTA(A2:A{nn})")]):
    ws.cell(nn + 2 + i, 1, lab).font = BOLD; ws.cell(nn + 2 + i, 2, f)
ws.cell(nn + 10, 1, "Примітка: strict — збіг точного (s,p,o). Recall занижений, бо експертна карта не є повним еталоном: частина LLM-триплетів справді є в тексті (див. Text-supported precision).")
# 6. Статистика карт
rows = [["Карта", "Понять", "Екземплярів", "Пропозицій", "Крос-зв'язків", "related_to, %"]]
st = dict(ev["map_stats"])
v2 = json.loads((ROOT / "models/cmap_v2.json").read_text(encoding="utf-8"))
import sys
sys.path.insert(0, str(ROOT / "scripts"))
from map_tools import stats   # noqa: E402
st["v2"] = stats(v2)
for k, s in st.items():
    rows.append([k, s["concepts"], s["instances"], s["propositions"], s["cross_links"], round(100 * s["related_to_share"], 1)])
sheet("Статистика карт", rows, [14, 10, 14, 14, 14, 14])
wb.save(ROOT / "docs/cmap_rubric.xlsx")
print("ok")
