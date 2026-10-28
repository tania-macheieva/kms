"""Тести ЛР6: каталог відношень, експертна карта, LLM-триплети, v2."""
import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from map_tools import REL, stats                 # noqa: E402
from check_triples import TEXT, norm             # noqa: E402
from validate_v2 import validate                 # noqa: E402


def jl(p):
    return [json.loads(l) for l in (ROOT / p).read_text(encoding="utf-8").splitlines() if l.strip()]


def test_catalog_size():
    assert 8 <= len(REL) <= 12 and all(r["definition"] for r in REL.values())


def test_expert_map_requirements():
    s = stats(json.loads((ROOT / "models/cmap_expert.json").read_text(encoding="utf-8")))
    assert s["concepts"] >= 25 and s["propositions"] >= 30 and s["cross_links"] >= 3 and s["instances"] >= 5
    assert s["related_to_share"] <= 0.10


def test_focus_doc():
    t = (ROOT / "docs/cmap_focus.md").read_text(encoding="utf-8")
    assert "?" in t and t.count("CQ") >= 5


def test_llm_triples():
    t = jl("models/triples_llm.jsonl")
    assert len(t) >= 40 and {x["source"] for x in t} >= {"gatk", "cdc"}
    found = sum(norm(x["evidence"]) in TEXT[x["source"]] for x in t)
    assert found / len(t) >= 0.9
    assert sum(1 for x in t if x["checks"]) >= 3


def test_v2_valid():
    m, rows, errs = validate()
    assert errs == []
    assert len(rows) >= 60
    assert all(r["p"] in REL for r in rows)
    assert set(r["status"] for r in rows) <= {"verified", "expert_asserted", "pending"}


def test_v2_changes_documented():
    t = (ROOT / "docs/cmap_changes_v1_v2.md").read_text(encoding="utf-8")
    assert t.count("| V2-") >= 5


def test_claims():
    r = list(csv.DictReader(open(ROOT / "claims/claims.csv", encoding="utf-8-sig")))
    assert r and all(x["status"] in {"verified", "rejected", "pending"} for x in r)
    assert all(x["verification_level"] in {"V0", "V1", "V2", "V3", "V4"} for x in r)


def test_rubric_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "docs/cmap_rubric.xlsx")
    assert {"Рубрика", "Оцінки", "Узгодженість", "Зауваження", "Precision-Recall"} <= set(wb.sheetnames)
    assert wb["Precision-Recall"].max_row >= 31
